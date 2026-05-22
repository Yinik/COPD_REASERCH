#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
将TSV审核文件转换为Excel (.xlsx) 格式
"""
import config  # 统一路径配置

import csv
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUTPUT_DIR = 'I:\\101实验专题\\关系抽取结果'

# 定义要转换的文件列表
FILES_TO_CONVERT = [
    ('待审核_精简版_explicit_高质量.tsv', '待审核_explicit_高质量.xlsx'),
    ('待审核_精简版_cooccur_兜底.tsv', '待审核_cooccur_兜底.xlsx'),
    ('待审核_精简版_全部_981条.tsv', '待审核_全部_981条.xlsx'),
]

def tsv_to_excel(tsv_path, xlsx_path):
    """将TSV文件转换为带格式的Excel文件"""
    
    # 读取TSV数据
    rows = []
    with open(tsv_path, 'r', encoding='utf-8-sig') as f:
        reader = csv.reader(f, delimiter='\t')
        header = next(reader)
        for row in reader:
            rows.append(row)
    
    # 创建Excel工作簿
    wb = Workbook()
    ws = wb.active
    ws.title = "三元组审核"
    
    # 定义样式
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # 设置列宽（根据字段名调整）
    col_widths = {
        '关系类型': 16,
        '头实体': 18,
        '头类型': 10,
        '尾实体': 18,
        '尾类型': 10,
        '头原文': 14,
        '尾原文': 14,
        '文献编号': 35,
        '匹配类型': 10,
        '精简上下文': 50,
        '审核状态': 10,
        '修改建议': 16,
        '备注': 20,
    }
    
    # 写入表头
    for col_idx, col_name in enumerate(header, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        # 设置列宽
        ws.column_dimensions[cell.column_letter].width = col_widths.get(col_name, 15)
    
    # 写入数据行
    for row_idx, row_data in enumerate(rows, 2):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            
            # 审核状态列添加颜色提示
            if header[col_idx - 1] == '审核状态' and value:
                if value == '保留':
                    cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                    cell.font = Font(color="006100")
                elif value == '删除':
                    cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                    cell.font = Font(color="9C0006")
                elif value == '修改':
                    cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
                    cell.font = Font(color="9C5700")
    
    # 冻结首行
    ws.freeze_panes = 'A2'
    
    # 设置行高
    ws.row_dimensions[1].height = 25
    for r in range(2, len(rows) + 2):
        ws.row_dimensions[r].height = 35
    
    # 添加数据验证（下拉列表）到审核状态列
    from openpyxl.worksheet.datavalidation import DataValidation
    
    status_col_idx = None
    for i, h in enumerate(header):
        if h == '审核状态':
            status_col_idx = i + 1
            break
    
    if status_col_idx:
        dv = DataValidation(type="list", formula1='"保留,删除,修改,待审核"', allow_blank=True)
        dv.error = '请从下拉列表中选择'
        dv.errorTitle = '输入错误'
        dv.prompt = '请选择审核状态'
        dv.promptTitle = '审核状态'
        
        last_row = len(rows) + 1
        dv.add(f'{ws.cell(row=2, column=status_col_idx).column_letter}2:{ws.cell(row=last_row, column=status_col_idx).column_letter}{last_row}')
        ws.add_data_validation(dv)
    
    # 保存
    wb.save(xlsx_path)
    print(f"  已生成: {os.path.basename(xlsx_path)} ({len(rows)} 条数据)")

# 执行转换
print("开始转换TSV → Excel...\n")
for tsv_name, xlsx_name in FILES_TO_CONVERT:
    tsv_path = os.path.join(OUTPUT_DIR, tsv_name)
    xlsx_path = os.path.join(OUTPUT_DIR, xlsx_name)
    
    if os.path.exists(tsv_path):
        tsv_to_excel(tsv_path, xlsx_path)
    else:
        print(f"  跳过: {tsv_name} (文件不存在)")

print("\n全部转换完成！")
