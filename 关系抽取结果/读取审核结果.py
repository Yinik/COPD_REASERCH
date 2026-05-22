#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
读取审核后的Excel文件，输出审核结果汇总和典型案例
"""
import config  # 统一路径配置

import os
from openpyxl import load_workbook

OUTPUT_DIR = 'I:\\101实验专题\\关系抽取结果'

def read_audit_file(filename):
    """读取审核Excel文件"""
    path = os.path.join(OUTPUT_DIR, filename)
    wb = load_workbook(path)
    ws = wb.active
    
    # 读取表头
    headers = [cell.value for cell in ws[1]]
    
    # 读取数据
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        row_dict = dict(zip(headers, row))
        rows.append(row_dict)
    
    return rows

# 读取Explicit
exp_rows = read_audit_file('医学审核_explicit_高质量.xlsx')
# 读取Cooccur
coo_rows = read_audit_file('医学审核_cooccur_兜底.xlsx')

# 分类
exp_kept = [r for r in exp_rows if r['审核状态'] == '保留']
exp_del = [r for r in exp_rows if r['审核状态'] == '删除']
coo_kept = [r for r in coo_rows if r['审核状态'] == '保留']
coo_del = [r for r in coo_rows if r['审核状态'] == '删除']

output_path = os.path.join(OUTPUT_DIR, '医学审核结果详细报告.txt')
with open(output_path, 'w', encoding='utf-8-sig') as f:
    f.write("=" * 80 + "\n")
    f.write("COPD医学文本关系抽取 — 基于医学知识的自动审核报告\n")
    f.write("=" * 80 + "\n\n")
    
    f.write("【审核总体统计】\n")
    f.write("-" * 60 + "\n")
    f.write(f"Explicit（明确模板匹配）:\n")
    f.write(f"  总数: {len(exp_rows)} 条\n")
    f.write(f"  保留: {len(exp_kept)} 条 ({len(exp_kept)/len(exp_rows)*100:.1f}%)\n")
    f.write(f"  删除: {len(exp_del)} 条 ({len(exp_del)/len(exp_rows)*100:.1f}%)\n\n")
    f.write(f"Cooccur（兜底共现匹配）:\n")
    f.write(f"  总数: {len(coo_rows)} 条\n")
    f.write(f"  保留: {len(coo_kept)} 条 ({len(coo_kept)/len(coo_rows)*100:.1f}%)\n")
    f.write(f"  删除: {len(coo_del)} 条 ({len(coo_del)/len(coo_rows)*100:.1f}%)\n\n")
    f.write(f"总计: {len(exp_rows)+len(coo_rows)} 条 → 保留 {len(exp_kept)+len(coo_kept)} 条 ({(len(exp_kept)+len(coo_kept))/(len(exp_rows)+len(coo_rows))*100:.1f}%)\n\n")
    
    # Explicit 保留案例
    f.write("=" * 80 + "\n")
    f.write("【Explicit 保留案例】（模板匹配质量高，医学上合理）\n")
    f.write("=" * 80 + "\n\n")
    
    # 按关系类型分组展示
    from collections import defaultdict
    exp_by_rel = defaultdict(list)
    for r in exp_kept:
        exp_by_rel[r['关系类型']].append(r)
    
    for rel_type in ['疾病-症状', '药物-治疗-疾病', '危险因素-疾病', '疾病-治疗', '检查-辅助诊断-疾病', '疾病-并发症']:
        items = exp_by_rel.get(rel_type, [])
        if not items:
            continue
        f.write(f"\n◆ {rel_type}（{len(items)} 条）\n")
        f.write("-" * 60 + "\n")
        for i, r in enumerate(items[:5], 1):
            f.write(f"  {i}. {r['头实体']} → {r['尾实体']}\n")
            f.write(f"     上下文: {r['精简上下文'][:100]}...\n")
            f.write(f"     审核依据: {r['备注']}\n\n")
        if len(items) > 5:
            f.write(f"     ... 等共 {len(items)} 条\n\n")
    
    # Explicit 删除案例
    if exp_del:
        f.write("\n" + "=" * 80 + "\n")
        f.write("【Explicit 删除案例】（医学常识错误或噪声实体）\n")
        f.write("=" * 80 + "\n\n")
        for i, r in enumerate(exp_del, 1):
            f.write(f"  {i}. ({r['关系类型']}) {r['头实体']} → {r['尾实体']}\n")
            f.write(f"     删除原因: {r['备注']}\n\n")
    
    # Cooccur 保留案例
    f.write("\n" + "=" * 80 + "\n")
    f.write("【Cooccur 保留案例】（共现但医学上确实相关）\n")
    f.write("=" * 80 + "\n\n")
    
    # 按是否有细化建议分组
    coo_with_suggestion = [r for r in coo_kept if r['修改建议']]
    coo_without_suggestion = [r for r in coo_kept if not r['修改建议']]
    
    if coo_with_suggestion:
        f.write(f"\n◆ 可细化为具体关系类型（{len(coo_with_suggestion)} 条）\n")
        f.write("-" * 60 + "\n")
        for i, r in enumerate(coo_with_suggestion[:15], 1):
            f.write(f"  {i}. ({r['头类型']}→{r['尾类型']}) {r['头实体']} → {r['尾实体']}\n")
            f.write(f"     建议改为: {r['修改建议']}\n")
            f.write(f"     审核依据: {r['备注']}\n")
            f.write(f"     上下文: {r['精简上下文'][:80]}...\n\n")
        if len(coo_with_suggestion) > 15:
            f.write(f"     ... 等共 {len(coo_with_suggestion)} 条\n\n")
    
    if coo_without_suggestion:
        f.write(f"\n◆ 保持为相关关系（{len(coo_without_suggestion)} 条）\n")
        f.write("-" * 60 + "\n")
        for i, r in enumerate(coo_without_suggestion[:10], 1):
            f.write(f"  {i}. ({r['头类型']}→{r['尾类型']}) {r['头实体']} → {r['尾实体']}\n")
            f.write(f"     审核依据: {r['备注']}\n\n")
    
    # Cooccur 删除案例（典型）
    f.write("\n" + "=" * 80 + "\n")
    f.write("【Cooccur 删除案例】（伪关系/方向错误/非核心关系）\n")
    f.write("=" * 80 + "\n\n")
    
    # 按删除原因分组
    del_reasons = defaultdict(list)
    for r in coo_del:
        reason = r['备注']
        del_reasons[reason].append(r)
    
    for reason, items in sorted(del_reasons.items(), key=lambda x: -len(x[1]))[:6]:
        f.write(f"\n◆ {reason}（{len(items)} 条）\n")
        f.write("-" * 60 + "\n")
        for i, r in enumerate(items[:3], 1):
            f.write(f"  {i}. ({r['头类型']}→{r['尾类型']}) {r['头实体']} → {r['尾实体']}\n")
            f.write(f"     上下文: {r['精简上下文'][:80]}...\n\n")
        if len(items) > 3:
            f.write(f"     ... 等共 {len(items)} 条\n\n")
    
    # 最终建议
    f.write("\n" + "=" * 80 + "\n")
    f.write("【审核结论与建议】\n")
    f.write("=" * 80 + "\n\n")
    total_kept = len(exp_kept) + len(coo_kept)
    total_del = len(exp_del) + len(coo_del)
    f.write(f"1. 审核后剩余 {total_kept} 条有效三元组（保留率 {(total_kept)/(total_kept+total_del)*100:.1f}%）\n")
    f.write(f"2. Explicit 质量极高，263条中仅删除2条（头痛不是COPD特异性症状）\n")
    f.write(f"3. Cooccur 经过严格筛选，718条中仅保留119条（16.6%），剔除了大量伪关系\n")
    f.write(f"4. 建议对保留的119条Cooccur中标注'修改建议'的进行细化\n")
    f.write(f"5. 最终可用于Neo4j导入的三元组约 {total_kept} 条\n\n")

print(f"报告已生成: {output_path}")
print(f"Explicit: 保留 {len(exp_kept)}, 删除 {len(exp_del)}")
print(f"Cooccur: 保留 {len(coo_kept)}, 删除 {len(coo_del)}")
